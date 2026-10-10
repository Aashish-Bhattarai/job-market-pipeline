# -*- coding: utf-8 -*-
"""
Imports pipeline/data/linkedin_applied.csv (LinkedIn's own "Applied" job
tracker, scraped read-only) into JobTracker.xlsx's Tracker sheet.

- Skips any LinkedIn row whose normalized (company, title) already exists
  anywhere in the Tracker sheet (rows 3+).
- Appends the rest after the last used row, with:
    Channel      = "LinkedIn (imported)"
    Date Applied = approximate date text derived from LinkedIn's relative
                   "Applied Xh/d/w/mo ago" string, relative to TODAY,
                   prefixed with "~" (e.g. "~12 Sep").
    Status       = "Viewed" if LinkedIn's status text mentions "viewed" or
                   "downloaded", else "Pending".
    Status Date  = left blank (LinkedIn does not give a separate timestamp
                   for the viewed/downloaded event).
    Notes        = "<linkedin status text> — <job url>"
  Fill colour per status copied from the existing STATUS_FILL convention
  used by update_tracker.py / update_tracker_2.py.
- Extends the Tracker autofilter to the new last row.
- Extends Dashboard!D6:D12 COUNTA/COUNTIF ranges from row 200 to row 1000,
  and the Status data-validation range on Tracker!E from row 500 to row 1000.

Safe to re-run: rows already present (by normalized company+title) are
skipped, so re-running after LinkedIn adds more "Applied" jobs will only
append the new ones.
"""
import csv
import datetime
import re
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

XLSX_PATH = "JobTracker.xlsx"
CSV_PATH = "pipeline/data/linkedin_applied.csv"
TODAY = datetime.datetime(2026, 9, 26)

STATUS_FILL = {
    "Rejected": "FFF8D7DA",
    "Interview": "FFD4EDDA",
    "Pending": "FFFFF3CD",
    "Viewed": "FFFFE8CC",
    "Incomplete": "FFE2E3E5",
}


def normalize(text):
    """Lowercase, drop parentheticals/punctuation, collapse whitespace,
    strip a few generic corporate suffixes, for loose duplicate matching."""
    if not text:
        return ""
    t = text.lower()
    t = re.sub(r"\([^)]*\)", " ", t)  # drop "(role not specified)" etc.
    t = re.sub(r"[^a-z0-9\s]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    for suffix in (" ltd", " limited", " plc", " llc", " inc", " corp",
                   " corporation", " group", " com"):
        if t.endswith(suffix):
            t = t[: -len(suffix)].strip()
    return t


def parse_applied_when(text):
    """'2w ago' -> TODAY - 14 days, etc. Returns a datetime (best-effort)."""
    m = re.search(r"(\d+)\s*(h|d|w|mo|y)\b", text or "")
    if not m:
        return TODAY
    n, unit = int(m.group(1)), m.group(2)
    if unit == "h":
        return TODAY  # sub-day precision not meaningful for a date field
    if unit == "d":
        return TODAY - datetime.timedelta(days=n)
    if unit == "w":
        return TODAY - datetime.timedelta(weeks=n)
    if unit == "mo":
        # approximate a month as 30 days
        return TODAY - datetime.timedelta(days=30 * n)
    if unit == "y":
        return TODAY - datetime.timedelta(days=365 * n)
    return TODAY


def approx_date_text(dt):
    return "~%d %s" % (dt.day, dt.strftime("%b"))


def main():
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        linkedin_rows = list(csv.DictReader(f))

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Tracker"]

    # Build the set of (normalized company, normalized title) already tracked.
    existing = set()
    last_used_row = 2
    r = 3
    while True:
        company = ws.cell(row=r, column=1).value
        role = ws.cell(row=r, column=2).value
        if company is None and role is None:
            break
        existing.add((normalize(company), normalize(role)))
        last_used_row = r
        r += 1

    scraped_count = len(linkedin_rows)
    already_present = []
    to_add = []
    for row in linkedin_rows:
        key = (normalize(row["company"]), normalize(row["title"]))
        if key in existing:
            already_present.append(row)
        else:
            to_add.append(row)
            existing.add(key)  # avoid double-adding duplicate LinkedIn rows against each other

    start_row = last_used_row + 1
    added_rows_summary = []
    for i, row in enumerate(to_add):
        r = start_row + i
        status_text = (row.get("linkedin_status") or "").strip()
        status = "Viewed" if re.search(r"viewed|downloaded", status_text, re.I) else "Pending"
        applied_dt = parse_applied_when(row.get("applied_when", ""))
        date_applied = approx_date_text(applied_dt)
        notes = f"{status_text} \u2014 {row['job_url']}" if status_text else row["job_url"]

        ws.cell(row=r, column=1, value=row["company"])
        ws.cell(row=r, column=2, value=row["title"])
        ws.cell(row=r, column=3, value="LinkedIn (imported)")
        ws.cell(row=r, column=4, value=date_applied)
        ws.cell(row=r, column=5, value=status)
        ws.cell(row=r, column=6, value=None)
        ws.cell(row=r, column=7, value=notes)

        fill = PatternFill(start_color=STATUS_FILL[status], end_color=STATUS_FILL[status], fill_type="solid")
        for col in range(1, 8):
            ws.cell(row=r, column=col).fill = fill

        added_rows_summary.append((r, row["company"], row["title"], status, date_applied))

    new_last_row = start_row + len(to_add) - 1 if to_add else last_used_row
    ws.auto_filter.ref = f"A2:G{new_last_row}"

    # Extend the Status data-validation range on Tracker!E from 500 to 1000.
    for dv in list(ws.data_validations.dataValidation):
        if str(dv.sqref).startswith("E3:E"):
            ws.data_validations.dataValidation.remove(dv)
            new_dv = DataValidation(type=dv.type, formula1=dv.formula1, allow_blank=dv.allow_blank)
            new_dv.add("E3:E1000")
            ws.add_data_validation(new_dv)

    # Extend Dashboard!D6:D12 formula ranges from row 200 to row 1000.
    dash = wb["Dashboard"]
    new_dashboard_formulas = {}
    for row_cells in dash.iter_rows():
        for cell in row_cells:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and ("Tracker!A3:A200" in v or "Tracker!E3:E200" in v):
                new_v = v.replace("Tracker!A3:A200", "Tracker!A3:A1000").replace("Tracker!E3:E200", "Tracker!E3:E1000")
                cell.value = new_v
                new_dashboard_formulas[cell.coordinate] = new_v

    wb.save(XLSX_PATH)

    print(f"LinkedIn scraped rows: {scraped_count}")
    print(f"Already present in Tracker (skipped): {len(already_present)}")
    print(f"Added to Tracker: {len(to_add)} (rows {start_row}-{new_last_row})")
    print("Sample added rows:")
    for r, company, title, status, date_applied in added_rows_summary[:5]:
        print(f"  row {r}: {company!r} | {title!r} | {status} | {date_applied}")
    print("New Dashboard formulas:")
    for coord, v in new_dashboard_formulas.items():
        print(f"  {coord}: {v}")
    print(f"Tracker autofilter: {ws.auto_filter.ref}")
    print("Status data validation ranges:", [str(dv.sqref) for dv in ws.data_validations.dataValidation])


if __name__ == "__main__":
    main()
