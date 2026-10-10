"""Render a tailored cover letter: pipeline/out/<job_id>/cover.json +
profile/facts.yaml -> Cover_Letter_<name>.docx + .pdf + .txt
(CONTRACTS.md #4).

Usage:
    python pipeline/render/render_cover.py <cover.json> <out_dir> [--facts <path>]
"""
from __future__ import annotations

import argparse
import html
import os
import sys

from docx import Document
from docx.shared import Pt, Inches

from common import (
    build_cover_content,
    convert_docx_to_pdf,
    convert_html_to_pdf_with_chrome,
    filename_stem,
    load_facts_bank,
    load_json,
    resolve_facts_path,
)

BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)


def _run(p, text, size=BODY_SIZE, bold=False):
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = size
    run.bold = bold
    return run


def build_docx(content: dict) -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _run(p, content["name"], bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    _run(p, content["contact_line"])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    _run(p, content["date"])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _run(p, "Dear Hiring Team,")

    for para_text in content["paragraphs"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        _run(p, para_text)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _run(p, "Kind regards,")
    p2 = doc.add_paragraph()
    _run(p2, content["name"])

    return doc


def build_txt(content: dict) -> str:
    lines = [
        content["name"],
        content["contact_line"],
        "",
        content["date"],
        "",
        "Dear Hiring Team,",
        "",
    ]
    for para_text in content["paragraphs"]:
        lines.append(para_text)
        lines.append("")
    lines.append("Kind regards,")
    lines.append(content["name"])
    return "\n".join(lines)


def build_html(content: dict) -> str:
    def esc(s: str) -> str:
        return html.escape(s or "")

    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<style>body{font-family:Calibri,Arial,sans-serif;font-size:11pt;",
        "margin:1in;} p{margin:0 0 10pt 0;}</style></head><body>",
        f"<p><b>{esc(content['name'])}</b><br>{esc(content['contact_line'])}</p>",
        f"<p>{esc(content['date'])}</p>",
        "<p>Dear Hiring Team,</p>",
    ]
    for para_text in content["paragraphs"]:
        parts.append(f"<p>{esc(para_text)}</p>")
    parts.append(f"<p>Kind regards,<br>{esc(content['name'])}</p>")
    parts.append("</body></html>")
    return "".join(parts)


def render(cover_path: str, out_dir: str, facts_path: str | None) -> dict:
    bank = load_facts_bank(resolve_facts_path(facts_path))
    cover = load_json(cover_path)
    content = build_cover_content(cover, bank)

    os.makedirs(out_dir, exist_ok=True)
    stem = filename_stem(content.get("name"))
    docx_path = os.path.join(out_dir, f"Cover_Letter_{stem}.docx")
    pdf_path = os.path.join(out_dir, f"Cover_Letter_{stem}.pdf")
    txt_path = os.path.join(out_dir, f"Cover_Letter_{stem}.txt")

    doc = build_docx(content)
    doc.save(docx_path)

    with open(txt_path, "w", encoding="utf-8") as fh:
        fh.write(build_txt(content))

    try:
        method = convert_docx_to_pdf(docx_path, pdf_path)
    except RuntimeError:
        html_path = os.path.join(out_dir, f"Cover_Letter_{stem}.html")
        with open(html_path, "w", encoding="utf-8") as fh:
            fh.write(build_html(content))
        method = convert_html_to_pdf_with_chrome(html_path, pdf_path)

    return {"docx": docx_path, "pdf": pdf_path, "txt": txt_path, "pdf_method": method}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("cover_path")
    parser.add_argument("out_dir")
    parser.add_argument("--facts", default=None)
    args = parser.parse_args()

    result = render(args.cover_path, args.out_dir, args.facts)
    print(f"DOCX: {result['docx']}")
    print(f"TXT:  {result['txt']}")
    print(f"PDF:  {result['pdf']} (via {result['pdf_method']})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
