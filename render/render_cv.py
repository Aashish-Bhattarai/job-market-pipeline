"""Render a tailored CV: pipeline/out/<job_id>/selection.json + profile/facts.yaml
-> CV_<name>.docx + .pdf (CONTRACTS.md #4).

Usage:
    python pipeline/render/render_cv.py <selection.json> <out_dir> [--facts <path>]

ATS-friendly rules: single column, no tables/text boxes/images, no header/
footer content, contact line lives in the document body, standard fonts
(Calibri 10.5-11pt) and standard section headings.
"""
from __future__ import annotations

import argparse
import html
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from common import (
    build_cv_content,
    convert_docx_to_pdf,
    convert_html_to_pdf_with_chrome,
    count_pdf_pages,
    filename_stem,
    load_facts_bank,
    load_json,
    resolve_facts_path,
    skill_group_label,
)

BODY_FONT = "Calibri"
BODY_SIZE = Pt(10.5)
NAME_SIZE = Pt(18)
HEADLINE_SIZE = Pt(12)
HEADING_SIZE = Pt(11)


def _set_run_font(run, size=BODY_SIZE, bold=False, italic=False):
    run.font.name = BODY_FONT
    run.font.size = size
    run.bold = bold
    run.italic = italic
    return run


def _add_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    _set_run_font(run, HEADING_SIZE, bold=True)
    # simple bottom border for a visual section rule without tables/text boxes
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    _set_run_font(run)
    return p


def _add_para(doc: Document, text: str, size=BODY_SIZE, bold=False, italic=False, space_after=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    _set_run_font(run, size, bold, italic)
    return p


def build_docx(content: dict) -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE

    # Name (body, not header)
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    name_p.paragraph_format.space_after = Pt(0)
    _set_run_font(name_p.add_run(content["name"]), NAME_SIZE, bold=True)

    if content.get("headline"):
        _add_para(doc, content["headline"], HEADLINE_SIZE, bold=False, space_after=Pt(2))

    _add_para(doc, content["contact_line"], BODY_SIZE, space_after=Pt(8))

    if content.get("summary"):
        _add_heading(doc, "Profile")
        _add_para(doc, content["summary"])

    if content.get("skills"):
        _add_heading(doc, "Technical Skills")
        for group, items in content["skills"].items():
            if not items:
                continue
            label = skill_group_label(group)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _set_run_font(p.add_run(f"{label}: "), BODY_SIZE, bold=True)
            _set_run_font(p.add_run(", ".join(items)), BODY_SIZE)

    if content.get("experience"):
        _add_heading(doc, "Professional Experience")
        for job in content["experience"]:
            header_bits = [b for b in [job.get("role"), job.get("org")] if b]
            _add_para(doc, " — ".join(header_bits), BODY_SIZE, bold=True, space_after=Pt(0))
            meta_bits = [b for b in [job.get("location"), job.get("dates")] if b]
            if meta_bits:
                _add_para(doc, " | ".join(meta_bits), BODY_SIZE, italic=True, space_after=Pt(2))
            for bullet in job.get("bullets", []):
                _add_bullet(doc, bullet)

    if content.get("projects"):
        _add_heading(doc, "Selected Projects")
        for proj in content["projects"]:
            header_bits = [b for b in [proj.get("role"), proj.get("org")] if b]
            title = " — ".join(header_bits) if proj.get("org") else (proj.get("role") or "")
            if title:
                _add_para(doc, title, BODY_SIZE, bold=True, space_after=Pt(0))
            meta_bits = [b for b in [proj.get("location"), proj.get("dates")] if b]
            if meta_bits:
                _add_para(doc, " | ".join(meta_bits), BODY_SIZE, italic=True, space_after=Pt(2))
            for bullet in proj.get("bullets", []):
                _add_bullet(doc, bullet)

    if content.get("education"):
        _add_heading(doc, "Education")
        for edu in content["education"]:
            header_bits = [b for b in [edu.get("degree"), edu.get("institution")] if b]
            _add_para(doc, " — ".join(header_bits), BODY_SIZE, bold=True, space_after=Pt(0))
            meta_bits = [b for b in [edu.get("location"), edu.get("dates")] if b]
            if meta_bits:
                _add_para(doc, " | ".join(meta_bits), BODY_SIZE, italic=True, space_after=Pt(2))
            for detail in edu.get("details", []) or []:
                _add_bullet(doc, detail)

    if content.get("certifications"):
        _add_heading(doc, "Certifications")
        for cert in content["certifications"]:
            year = cert.get("year")
            text = f"{cert.get('name')} ({year})" if year else cert.get("name", "")
            _add_bullet(doc, text)

    _add_heading(doc, "Right to Work")
    _add_para(doc, content["right_to_work"])

    return doc


def build_html(content: dict) -> str:
    def esc(s: str) -> str:
        return html.escape(s or "")

    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<style>",
        "body{font-family:Calibri,Arial,sans-serif;font-size:10.5pt;margin:0.5in 0.7in;}",
        "h1{font-size:18pt;margin:0 0 2pt 0;}",
        ".headline{font-size:12pt;margin:0 0 2pt 0;}",
        ".contact{margin:0 0 8pt 0;}",
        "h2{font-size:11pt;text-transform:uppercase;border-bottom:1px solid #000;",
        "margin:10pt 0 2pt 0;padding-bottom:1pt;}",
        "p{margin:0 0 4pt 0;}",
        "ul{margin:0 0 4pt 0;padding-left:18pt;}",
        "li{margin:0 0 2pt 0;}",
        ".meta{font-style:italic;margin:0 0 2pt 0;}",
        ".role{font-weight:bold;margin:0;}",
        "</style></head><body>",
        f"<h1>{esc(content['name'])}</h1>",
    ]
    if content.get("headline"):
        parts.append(f"<div class='headline'>{esc(content['headline'])}</div>")
    parts.append(f"<div class='contact'>{esc(content['contact_line'])}</div>")

    if content.get("summary"):
        parts.append("<h2>Profile</h2>")
        parts.append(f"<p>{esc(content['summary'])}</p>")

    if content.get("skills"):
        parts.append("<h2>Technical Skills</h2>")
        for group, items in content["skills"].items():
            if not items:
                continue
            label = skill_group_label(group)
            parts.append(f"<p><b>{esc(label)}:</b> {esc(', '.join(items))}</p>")

    def render_block(title, entries, heading):
        parts.append(f"<h2>{heading}</h2>")
        for e in entries:
            header_bits = [b for b in [e.get("role"), e.get("org")] if b]
            if header_bits:
                parts.append(f"<p class='role'>{esc(' — '.join(header_bits))}</p>")
            meta_bits = [b for b in [e.get("location"), e.get("dates")] if b]
            if meta_bits:
                parts.append(f"<p class='meta'>{esc(' | '.join(meta_bits))}</p>")
            bullets = e.get("bullets") or e.get("details") or []
            if bullets:
                parts.append("<ul>" + "".join(f"<li>{esc(b)}</li>" for b in bullets) + "</ul>")

    if content.get("experience"):
        render_block("experience", content["experience"], "Professional Experience")
    if content.get("projects"):
        render_block("projects", content["projects"], "Selected Projects")
    if content.get("education"):
        render_block("education", content["education"], "Education")
        # education uses "details" not "bullets"; render_block already checks both

    if content.get("certifications"):
        parts.append("<h2>Certifications</h2><ul>")
        for cert in content["certifications"]:
            year = cert.get("year")
            text = f"{cert.get('name')} ({year})" if year else cert.get("name", "")
            parts.append(f"<li>{esc(text)}</li>")
        parts.append("</ul>")

    parts.append("<h2>Right to Work</h2>")
    parts.append(f"<p>{esc(content['right_to_work'])}</p>")
    parts.append("</body></html>")
    return "".join(parts)


def render(selection_path: str, out_dir: str, facts_path: str | None) -> dict:
    bank = load_facts_bank(resolve_facts_path(facts_path))
    selection = load_json(selection_path)
    content = build_cv_content(selection, bank)

    os.makedirs(out_dir, exist_ok=True)
    stem = filename_stem(content.get("name"))
    docx_path = os.path.join(out_dir, f"CV_{stem}.docx")
    pdf_path = os.path.join(out_dir, f"CV_{stem}.pdf")

    doc = build_docx(content)
    doc.save(docx_path)

    try:
        method = convert_docx_to_pdf(docx_path, pdf_path)
    except RuntimeError:
        html_path = os.path.join(out_dir, f"CV_{stem}.html")
        with open(html_path, "w", encoding="utf-8") as fh:
            fh.write(build_html(content))
        method = convert_html_to_pdf_with_chrome(html_path, pdf_path)

    page_count = count_pdf_pages(pdf_path)
    return {"docx": docx_path, "pdf": pdf_path, "pdf_method": method, "page_count": page_count}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("selection_path")
    parser.add_argument("out_dir")
    parser.add_argument("--facts", default=None)
    parser.add_argument(
        "--max-pages",
        type=int,
        default=None,
        help="Fail (non-zero exit) if the rendered PDF exceeds this many pages.",
    )
    args = parser.parse_args()

    result = render(args.selection_path, args.out_dir, args.facts)
    print(f"DOCX: {result['docx']}")
    print(f"PDF:  {result['pdf']} (via {result['pdf_method']}, {result['page_count']} page(s))")

    if args.max_pages is not None and result["page_count"] > args.max_pages:
        print(
            f"FAIL: CV PDF is {result['page_count']} page(s), exceeds --max-pages {args.max_pages}",
            file=sys.stderr,
        )
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
